# -*- coding: utf-8 -*-
"""Build a "cas non reussis" review workbook for the OKILYS website, from the last audit run.

    py _notes/tests/build_site_review.py

Lists every case that is NOT Pass (Fail + Blocked) with what it tests, why it did not pass and who must act.
Output: _notes/validation/cas-non-reussis-site-web-<version>.xlsx
"""
import glob
import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
VAL = os.path.join(ROOT, "_notes", "validation")
NAVY, TEAL = "1A2333", "21778A"
FILL = {"Fail": "F8CBAD", "Blocked": "FFF2CC"}
STATUT_FR = {"Fail": "Échec (défaut)", "Blocked": "Bloqué"}
AREA_FR = {"UT": "Test unitaire", "IT": "Test d'intégration", "ST": "Test système (navigateur)", "PT": "Test d'intrusion"}


def titles():
    runs = sorted(glob.glob(os.path.join(VAL, "site-spec-run-*.md")))
    out = {}
    if runs:
        for m in re.finditer(r"^\|\s*((?:UT|IT|ST|PT)-[A-Z0-9]+-\d+)\s*\|\s*(.*?)\s*\|", open(runs[-1], encoding="utf-8").read(), re.M):
            out.setdefault(m.group(1), m.group(2))
    return out


def spec_ids():
    """The case ids actually defined in the specification (definition tables), so the review lists only real cases."""
    spec = open(os.path.join(ROOT, "_notes", "TEST-SPECIFICATION-site-web.md"), encoding="utf-8").read()
    return {m.group(1) for m in re.finditer(r"^\|\s*((?:UT|IT|ST|PT)-[A-Z0-9]+-\d+)\s*\|", spec, re.M)}


def who(cid, status, note):
    n = (note or "").lower()
    if status == "Fail":
        return "Défaut - voir le classeur des défauts (à corriger ou non)"
    if any(w in n for w in ("web3forms", "registrar", "dnssec", "dépôt", "depot", "github", "accès", "acces", "accord", "autorises", "envoi réel", "envoi reel")):
        return "Toi (accès ou accord requis)"
    if cid.startswith("ST"):
        return "Testeur (scénario navigateur : Chrome / Edge / Firefox / Safari iOS)"
    return "Testeur / navigateur"


def build():
    with open(os.path.join(VAL, "site-spec-results.json"), encoding="utf-8") as f:
        data = json.load(f)
    version = data.get("version", "site")
    T = titles()
    ids = spec_ids()
    defects = data.get("defects", {})
    seen = set()
    bad = [r for r in data["records"] if r["status"] in ("Fail", "Blocked") and r["id"] in ids and not (r["id"] in seen or seen.add(r["id"]))]
    order = {"Fail": 0, "Blocked": 1}
    bad.sort(key=lambda r: (order.get(r["status"], 2), r["id"]))
    wb = Workbook()
    ws = wb.active
    ws.title = "Cas non reussis"
    ws["A1"] = f"OKILYS site web (www.okilys.com) - Cas de test non reussis (version {version}) - a revoir"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    n_fail = sum(1 for r in bad if r["status"] == "Fail")
    n_block = sum(1 for r in bad if r["status"] == "Blocked")
    ws["A2"] = (f"{len(bad)} cas non reussis : {n_fail} echecs (defauts - detail dans defauts-site-web) et {n_block} bloques. "
                "Un cas 'Bloque' n'est pas un defaut : c'est un test qui n'a pas pu etre joue - il faut soit un testeur dans un navigateur (scenarios ST), "
                "soit tes acces (envoi reel du formulaire, tableau Web3Forms, registrar, depot GitHub). La colonne 'Qui doit agir' precise quoi faire.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 64
    heads = ["Id", "Niveau", "Statut", "Ce que le cas teste", "Pourquoi il n'est pas passe / ce qu'il faut", "Qui doit agir"]
    ws.append([])
    ws.append(heads)
    for c in ws[4]:
        c.fill, c.font = PatternFill("solid", fgColor=TEAL), Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in bad:
        cid, status = r["id"], r["status"]
        d = defects.get(cid, {}) or {}
        what = T.get(cid, "")
        reason = (d.get("plain_fr") if status == "Fail" and d else None) or r.get("message", "").replace("BLOCKED: ", "")
        ws.append([cid, AREA_FR.get(cid[:2], ""), STATUT_FR.get(status, status), what, reason[:400], who(cid, status, r.get("message"))])
        row = ws.max_row
        for c in ws[row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=3).fill = PatternFill("solid", fgColor=FILL.get(status, "FFFFFF"))
    for col, w in zip("ABCDEF", (13, 22, 16, 50, 62, 46)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:F{ws.max_row}"
    path = os.path.join(VAL, f"cas-non-reussis-site-web-{version}.xlsx")
    wb.save(path)
    print(f"Workbook: {path} ({len(bad)} cas : {n_fail} echecs, {n_block} bloques)")
    return path


if __name__ == "__main__":
    build()
