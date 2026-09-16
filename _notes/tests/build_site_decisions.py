# -*- coding: utf-8 -*-
"""Build the defect-decision workbook for the OKILYS website, from the last audit run.

    py _notes/tests/build_site_decisions.py

Input : _notes/validation/site-spec-results.json (produced by site_audit.py).
Output: _notes/validation/defauts-site-web-<version>.xlsx  - one row per open defect (status Fail), same layout and
        columns as the CTMS decision workbook (docs/validation/defauts-decisions-*.xlsx) so both compile together:
        Id / Gravité / Domaine / Où / Ce qui ne va pas / Impact / Traduction technique (FR) / Emplacement /
        Action souhaitée (dropdown) / Vos notes, plus a "Mode d'emploi" sheet.
"""
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
VAL = os.path.join(ROOT, "_notes", "validation")
SEV_RANK = {"Critique": 5, "Critical": 5, "Haut": 4, "High": 4, "Majeur": 4, "Major": 4, "Moyen": 3, "Medium": 3, "Mineur": 2, "Minor": 2, "Faible": 1, "Low": 1, "": 0}
SEV_FR = {"Critical": "Critique", "High": "Haut", "Major": "Majeur", "Medium": "Moyen", "Minor": "Mineur", "Low": "Faible"}
AREA_FR = {"UT": "Test unitaire", "IT": "Test d'intégration", "ST": "Test système", "PT": "Test d'intrusion"}
ACTIONS = '"Corriger,Corriger en priorite,Corriger plus tard,Ne pas corriger,A decider"'
NAVY, TEAL, YELLOW = "1A2333", "21778A", "FFF2CC"


def build():
    with open(os.path.join(VAL, "site-spec-results.json"), encoding="utf-8") as f:
        data = json.load(f)
    version = data.get("version", "site")
    by_id = {r["id"]: r for r in data["records"]}
    defects = data.get("defects", {})
    fails = sorted((cid for cid, r in by_id.items() if r["status"] == "Fail"),
                   key=lambda c: -SEV_RANK.get(SEV_FR.get((defects.get(c, {}) or {}).get("severity", ""), (defects.get(c, {}) or {}).get("severity", "")), 0))

    wb = Workbook()
    ws = wb.active
    ws.title = "Defauts - decisions"
    ws["A1"] = f"OKILYS site web (www.okilys.com) - Registre des defauts (version {version}) : vos decisions"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = (f"Ecrivez votre choix dans la colonne Action souhaitee (menu deroulant) et, si besoin, une precision dans Vos notes. {len(fails)} defauts, tous Mineurs, tries du plus grave au plus leger. "
                "Meme format que le classeur du CTMS (defauts-decisions) pour compiler les deux ensemble. Colonnes : Ou = la page/section/encart ; Ce qui ne va pas = en langage simple ; "
                "Impact = ce que ca change pour le visiteur ou le referencement ; Traduction technique = version detaillee FR ; Emplacement = fichier. Cellules a remplir = fond jaune. "
                "La session de tests ne corrige pas le site : elle signale ; la session Site Web corrige.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:J2")
    ws.row_dimensions[2].height = 70
    heads = ["Id", "Gravite", "Domaine", "Ou (page / section / encart)", "Ce qui ne va pas (langage simple)", "Impact si non corrige", "Traduction technique (FR)", "Emplacement (fichier)", "Action souhaitee", "Vos notes"]
    ws.append([])
    ws.append(heads)
    for c in ws[4]:
        c.fill, c.font = PatternFill("solid", fgColor=TEAL), Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for cid in fails:
        d = defects.get(cid, {}) or {}
        grav = SEV_FR.get(d.get("severity", ""), d.get("severity") or "?")
        ws.append([cid, grav, AREA_FR.get(cid[:2], ""), d.get("ref_fr") or "", d.get("plain_fr") or by_id[cid]["message"], d.get("impact_fr") or "", d.get("defect_fr") or "", d.get("location") or "", None, None])
        row = ws.max_row
        for c in ws[row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=9).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(row=row, column=10).fill = PatternFill("solid", fgColor=YELLOW)
    last = ws.max_row
    if last >= 5:
        dv = DataValidation(type="list", formula1=ACTIONS, allow_blank=True)
        dv.add(f"I5:I{last}")
        ws.add_data_validation(dv)
    for col, w in zip("ABCDEFGHIJ", (12, 12, 18, 34, 52, 40, 50, 26, 22, 26)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:J{max(last, 4)}"

    ws2 = wb.create_sheet("Mode d'emploi")
    ws2.column_dimensions["A"].width = 110
    lines = ["Comment ca marche",
             "1) Dans l'onglet Defauts - decisions, remplissez la colonne Action souhaitee pour chaque ligne (menu deroulant).",
             "2) Renvoyez-le a la session Site Web (ou dites simplement les Id a corriger) : c'est elle qui modifie le site.",
             "3) Apres correction et publication, la session de tests relance py _notes/tests/site_audit.py : le cas repasse de Echec a Reussite.",
             None,
             f"Volumes (site www.okilys.com, version {version}) - {len(fails)} defauts ouverts, tous Mineurs :",
             "- UT-PAGE-02 : titre de suite.html trop long (72 caracteres).",
             "- UT-PAGE-03 : 3 descriptions un peu longues (161-162 caracteres).",
             "- UT-PAGE-16 : logos partenaires et quelques images sans dimensions (width/height).",
             "- IT-NAV-03 : libelle du pied de page different du menu (Fondatrice).",
             "- PT-HDR-01 : pas d'en-tetes de securite (limite GitHub Pages, connu).",
             "- PT-DNS-01 : pas d'enregistrement DMARC sur okilys.com (a ajouter chez le registrar).",
             None,
             "Bloques (a traiter avec vos acces) : envoi reel du formulaire de contact, analyse Web3Forms, DNSSEC registrar, scan du depot GitHub, scenarios navigateur. Voir le rapport site-spec-run-*.md.",
             "Aucun defaut Critique ni Majeur. Le rapport complet et les resultats bruts sont dans _notes/validation/."]
    for i, t in enumerate(lines, 1):
        if t is not None:
            ws2.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical="top")
    ws2["A1"].font = Font(bold=True, size=13, color=NAVY)

    path = os.path.join(VAL, f"defauts-site-web-{version}.xlsx")
    wb.save(path)
    print(f"Workbook: {path} ({len(fails)} defects)")
    return path


if __name__ == "__main__":
    build()
